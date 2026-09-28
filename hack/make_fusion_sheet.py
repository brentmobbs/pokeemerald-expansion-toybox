# Builds docs/hack/fusion_stats.xlsx: base and fusion stats per the rules in docs/hack/DESIGN.md.
# Base stats come from expansion's species data (latest-generation values).
import itertools
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

BASES = [  # name, type (group), original gen, HP, Atk, Def, SpA, SpD, Spe
    ("Porygon","Normal",1,65,60,70,85,75,40), ("Dunsparce","Normal",2,100,70,70,65,65,45), ("Teddiursa","Normal",2,60,80,50,50,50,40),
    ("Growlithe","Fire",1,55,70,45,70,50,60), ("Cyndaquil","Fire",2,39,52,43,60,50,65), ("Numel","Fire",3,60,60,40,65,45,35),
    ("Azurill","Water",3,50,20,40,20,40,20), ("Corsola","Water",2,65,55,95,65,95,35), ("Mudkip","Water",3,50,70,50,50,50,40),
    ("Pichu","Electric",2,20,40,15,35,35,60), ("Voltorb","Electric",1,40,30,50,55,55,100), ("Mareep","Electric",2,55,40,40,65,45,35),
    ("Bulbasaur","Grass",1,45,49,49,65,65,45), ("Celebi","Grass",2,100,100,100,100,100,100), ("Shroomish","Grass",3,60,40,60,40,60,35),
    ("Swinub","Ice",2,50,50,40,30,30,50), ("Snorunt","Ice",3,50,50,50,50,50,50), ("Spheal","Ice",3,70,40,50,55,50,25),
    ("Mankey","Fighting",1,40,80,35,35,45,70), ("Tyrogue","Fighting",2,35,35,35,35,35,35), ("Makuhita","Fighting",3,72,60,30,20,30,25),
    ("Ekans","Poison",1,35,60,44,40,54,55), ("Gulpin","Poison",3,70,43,53,43,53,40), ("Koffing","Poison",1,40,65,95,60,45,35),
    ("Sandshrew","Ground",1,50,75,85,20,30,40), ("Phanpy","Ground",2,90,60,60,40,40,40), ("Baltoy","Ground",3,40,40,55,40,70,55),
    ("Zubat","Flying",1,40,45,35,30,40,55), ("Murkrow","Flying",2,60,85,42,85,42,91), ("Swablu","Flying",3,45,40,60,40,75,50),
    ("Mew","Psychic",1,100,100,100,100,100,100), ("Spoink","Psychic",3,60,25,35,70,80,60), ("Jirachi","Psychic",3,100,100,100,100,100,100),
    ("Paras","Bug",1,35,70,55,45,55,25), ("Pineco","Bug",2,50,65,90,35,35,15), ("Wurmple","Bug",3,45,45,35,20,30,20),
    ("Shuckle","Rock",2,20,10,230,10,230,5), ("Larvitar","Rock",2,50,64,50,45,50,41), ("Nosepass","Rock",3,30,45,135,45,90,30),
    ("Gastly","Ghost",1,30,35,30,100,35,80), ("Misdreavus","Ghost",2,60,60,60,85,85,85), ("Sableye","Ghost",3,50,75,75,65,65,50),
    ("Dratini","Dragon",1,41,64,45,50,50,50), ("Bagon","Dragon",3,45,75,60,40,30,50),
    ("Sneasel","Dark",2,55,95,55,35,75,115), ("Houndour","Dark",2,45,60,30,80,50,65), ("Carvanha","Dark",3,45,90,20,65,20,65),
    ("Magnemite","Steel",1,25,35,70,95,55,45), ("Aron","Steel",3,50,70,100,40,40,30), ("Beldum","Steel",3,40,55,80,35,60,30),
    ("Cleffa","Fairy",2,50,25,28,45,55,15), ("Togepi","Fairy",2,35,20,65,40,65,20),
]
STATS = ["HP","Atk","Def","SpA","SpD","Spe"]
F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF"); GREEN = Font(name="Arial", size=10, color="008000")
HEAD = PatternFill("solid", start_color="DDDDDD"); INPUT = PatternFill("solid", start_color="FFFF00")

def header(ws, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h); c.font = FB; c.fill = HEAD; c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"

def scaled_block(ws, r, raw_cols, target, final_col, help_col):
    """Scale raw stats to exactly `target` total (largest remainder method).
    Writes final stats at final_col.., helpers from help_col.."""
    raw = [f"{L(c)}{r}" for c in raw_cols]
    rawsum = f"SUM({raw[0]}:{raw[-1]})"
    ex = [help_col + i for i in range(6)]; fl = [help_col + 6 + i for i in range(6)]
    left = help_col + 12; rem = [help_col + 13 + i for i in range(6)]
    for i in range(6):
        ws.cell(row=r, column=ex[i], value=f"={raw[i]}*{target}/{rawsum}")
        ws.cell(row=r, column=fl[i], value=f"=INT({L(ex[i])}{r})")
        ws.cell(row=r, column=rem[i], value=f"=ROUND({L(ex[i])}{r}-{L(fl[i])}{r},9)")
    ws.cell(row=r, column=left, value=f"={target}-SUM({L(fl[0])}{r}:{L(fl[-1])}{r})")
    remrng = f"${L(rem[0])}{r}:${L(rem[-1])}{r}"
    for i in range(6):
        me = f"{L(rem[i])}{r}"
        rank = f"SUMPRODUCT(--({remrng}>{me}))"
        if i > 0:
            rank += f"+SUMPRODUCT(--(${L(rem[0])}{r}:{L(rem[i-1])}{r}={me}))"
        ws.cell(row=r, column=final_col + i, value=f"={L(fl[i])}{r}+IF({rank}<${L(left)}{r},1,0)")
    ws.cell(row=r, column=final_col + 6, value=f"=SUM({L(final_col)}{r}:{L(final_col+5)}{r})")
    return help_col + 19

HELP_HEAD = [f"{s} exact" for s in STATS] + [f"{s} floor" for s in STATS] + ["Points left"] + [f"{s} remainder" for s in STATS]

wb = Workbook()
# ---- Settings
st = wb.active; st.title = "Settings"
rows = [("Setting", "Value", "Note"),
        ("Base BST", 400, "Owner's rule: every base is scaled to this total."),
        ("Fusion BST", 500, "Owner's rule: every fusion is scaled to this total."),
        ("Weight of higher stat", "=2/3", "Owner's rule: fusion stat = 2/3 of the higher parent + 1/3 of the lower, per stat."),
        ("Weight of lower stat", "=1-B4", ""),
        ("", "", ""),
        ("How rounding works", "", "Stats are scaled, then rounded down; the points still missing from the target total go one each to the stats that lost the most in rounding (ties: earlier stat first). This keeps every total exact."),
        ("Fusion stats use", "", "The bases' 400-scaled stats (owner's rule), not their original stats. The 2/3 + 1/3 mix is not rounded before scaling to 500."),
        ("Original stats source", "", "pokeemerald-expansion species data, latest-generation values."),
        ("Yellow cells", "", "Inputs you can change. Everything else is calculated."),
        ("Placeholder names", "", "Fusion names are auto-made placeholders (first half of A + second half of B)."),]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row, 1):
        cell = st.cell(row=r, column=c, value=v); cell.font = FB if r == 1 else F
for r in (2, 3): st.cell(row=r, column=2).font = BLUE; st.cell(row=r, column=2).fill = INPUT
st.column_dimensions["A"].width = 24; st.column_dimensions["B"].width = 10; st.column_dimensions["C"].width = 110

# ---- Bases
b = wb.create_sheet("Bases")
cols = ["ID", "Name", "Type", "Orig. gen", "Ability"] + [f"Orig {s}" for s in STATS] + ["Orig BST"] + STATS + ["BST"] + HELP_HEAD
header(b, cols)
for i, (name, typ, gen, *stats) in enumerate(BASES, 1):
    r = i + 1
    for c, v in enumerate([i, name, typ, gen, None] + stats, 1):
        cell = b.cell(row=r, column=c, value=v); cell.font = BLUE if c >= 6 else F
    b.cell(row=r, column=5).fill = INPUT
    b.cell(row=r, column=12, value=f"=SUM(F{r}:K{r})").font = F
    scaled_block(b, r, range(6, 12), "Settings!$B$2", 13, 20)
b.cell(row=2, column=5).comment = Comment("Fill in each base's one ability here. Fusions pick it up automatically.", "Claude")
for c in range(1, 20): b.column_dimensions[L(c)].width = 9
b.column_dimensions["B"].width = 12; b.column_dimensions["E"].width = 16
b.column_dimensions.group("T", L(19 + 19), hidden=True)

# ---- Fusions
fz = wb.create_sheet("Fusions")
cols = ["#", "Parent A ID", "Parent B ID", "Parent A", "Parent B", "Name (placeholder)", "Type 1", "Type 2",
        "Ability 1 (A)", "Ability 2 (B)", "Ability 3 (type pool)"] + STATS + ["BST"] + [f"Mix {s}" for s in STATS] + HELP_HEAD
header(fz, cols)
bref = lambda col, idc, r: f"INDEX(Bases!${col}$2:${col}$53,{idc}{r})"
for n, (a, bb) in enumerate(itertools.combinations(range(1, len(BASES) + 1), 2), 1):
    r = n + 1
    fz.cell(row=r, column=1, value=n); fz.cell(row=r, column=2, value=a); fz.cell(row=r, column=3, value=bb)
    fz.cell(row=r, column=4, value="=" + bref("B", "B", r)); fz.cell(row=r, column=5, value="=" + bref("B", "C", r))
    fz.cell(row=r, column=6, value=f"=LEFT(D{r},ROUNDUP(LEN(D{r})/2,0))&LOWER(MID(E{r},ROUNDUP(LEN(E{r})/2,0)+1,20))")
    fz.cell(row=r, column=7, value="=" + bref("C", "B", r))
    fz.cell(row=r, column=8, value=f'=IF({bref("C","C",r)}=G{r},"",{bref("C","C",r)})')
    fz.cell(row=r, column=9, value=f'=IF({bref("E","B",r)}="","",{bref("E","B",r)})')
    fz.cell(row=r, column=10, value=f'=IF({bref("E","C",r)}="","",{bref("E","C",r)})')
    fz.cell(row=r, column=11).fill = INPUT
    for i in range(6):  # mix of the parents' 400-scaled stats
        col = L(13 + i)
        pa, pb = bref(col, "B", r), bref(col, "C", r)
        fz.cell(row=r, column=19 + i, value=f"=Settings!$B$4*MAX({pa},{pb})+Settings!$B$5*MIN({pa},{pb})")
    scaled_block(fz, r, range(19, 25), "Settings!$B$3", 12, 25)
    for c in range(1, 25 + 19):
        cell = fz.cell(row=r, column=c); cell.font = BLUE if c in (2, 3) else (GREEN if c in (4, 5, 7, 8, 9, 10) else F)
for c in range(1, 25): fz.column_dimensions[L(c)].width = 9
for c, w in (("D", 12), ("E", 12), ("F", 16), ("I", 14), ("J", 14), ("K", 18)): fz.column_dimensions[c].width = w
fz.column_dimensions.group("Y", L(24 + 19), hidden=True)
fz.auto_filter.ref = f"A1:{L(24 + 19)}{len(BASES) * (len(BASES) - 1) // 2 + 1}"

wb.save("docs/hack/fusion_stats.xlsx")
print("saved")
